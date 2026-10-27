"""Development validation reports and original-row error transitions."""
from pathlib import Path
import html
import os

import numpy as np
import pandas as pd

from .protocol import ACTIVE, metrics, choose_threshold, choose_best, pair_weights


def csv(path, frame):
    path = Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    tmp = path.with_name(path.name+'.tmp')
    frame.to_csv(tmp,index=False,encoding='utf-8-sig')
    os.replace(tmp,path)


def flat_trials(records):
    rows = []
    for r in records:
        w1,w3 = pair_weights(r['indices'])
        for policy,m in r['validation_metrics'].items():
            rows.append(dict(trial_id=r['trial_id'],trial_number=r['trial_number'],
                phase=r['proposal']['method'],V01_weight=w1,V03_weight=w3,
                split='val',threshold_policy=policy,best_iteration=r['best_iteration'],
                structure_weight_sum_positive=r['weight_summary']['positive_sum'],
                structure_weight_sum_negative=r['weight_summary']['negative_sum'],
                fit_seconds=r['fit_seconds'],**m))
    return pd.DataFrame(rows)


def progress(records):
    rows=[];best=-np.inf
    for r in records:
        auc=r['validation_metrics']['fixed_0p5']['AUROC'];best=max(best,auc)
        rows.append(dict(completed_unique_configurations=r['trial_number'],
            validation_AUROC=auc,best_validation_AUROC_so_far=best))
    return pd.DataFrame(rows)


def progress_svg(path, records):
    if not records:return
    f=progress(records);n=len(f);vals=f.validation_AUROC.to_numpy()
    lo,hi=float(vals.min()),float(vals.max());pad=max((hi-lo)*.15,.0005)
    lo-=pad;hi+=pad
    W,H=1100,650;left,top,right,bottom=100,85,1040,550
    xx=lambda t:left+(right-left)*(t-1)/max(n-1,1)
    yy=lambda t:bottom-(t-lo)/(hi-lo)*(bottom-top)
    parts=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}">',
      '<rect width="100%" height="100%" fill="white"/>',
      '<g font-family="Arial, sans-serif" fill="#243B53">',
      '<text x="100" y="36" font-size="25" font-weight="bold">V01 and V03 Bayesian weight search</text>',
      '<text x="100" y="63" font-size="16">Validation development AUROC; 1.00-3.00 in increments of 0.05</text>']
    for y in np.linspace(lo,hi,6):
        Y=yy(y);parts.extend([f'<line x1="{left}" y1="{Y:.2f}" x2="{right}" y2="{Y:.2f}" stroke="#DCE3EA"/>',
            f'<text x="{left-12}" y="{Y+5:.2f}" text-anchor="end" font-size="15">{y:.5f}</text>'])
    for x in sorted(set([1,n]+[int(v) for v in np.linspace(1,n,6)])):
        X=xx(x);parts.append(f'<text x="{X:.2f}" y="{bottom+27}" text-anchor="middle" font-size="15">{x}</text>')
    baseline=float(vals[0]);Y=yy(baseline)
    parts.append(f'<line x1="{left}" y1="{Y:.2f}" x2="{right}" y2="{Y:.2f}" stroke="#E29532" stroke-width="2" stroke-dasharray="7 5"/>')
    for x,y in zip(f.completed_unique_configurations,vals):
        parts.append(f'<circle cx="{xx(x):.2f}" cy="{yy(y):.2f}" r="3" fill="#93A8BA"/>')
    pts=' '.join(f'{xx(x):.2f},{yy(y):.2f}' for x,y in zip(f.completed_unique_configurations,f.best_validation_AUROC_so_far))
    parts.append(f'<polyline points="{pts}" fill="none" stroke="#2463A8" stroke-width="3"/>')
    parts.extend(['<text x="570" y="610" text-anchor="middle" font-size="18">Completed unique weight combinations</text>',
        '<text transform="translate(24,320) rotate(-90)" text-anchor="middle" font-size="18">Validation AUROC</text>',
        '<text x="100" y="638" font-size="14">Blue: best so far | Gray: evaluated combinations | Orange dashed: unweighted control</text>',
        '</g></svg>'])
    Path(path).write_text('\n'.join(parts),encoding='utf-8')


def workbook(path, frames):
    from openpyxl.styles import Font, PatternFill, Alignment
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    tmp=path.with_name(path.stem+'.tmp.xlsx')
    with pd.ExcelWriter(tmp,engine='openpyxl') as writer:
        for name,f in frames.items():
            f.to_excel(writer,sheet_name=name,index=False)
            ws=writer.sheets[name];ws.freeze_panes='A2'
            if ws.max_column:ws.auto_filter.ref=ws.dimensions
            for c in ws[1]:
                c.font=Font(name='Arial',bold=True,color='FFFFFF')
                c.fill=PatternFill('solid',fgColor='243B53');c.alignment=Alignment(wrap_text=True,vertical='center')
            ws.row_dimensions[1].height=30
            for col in ws.columns:
                vals=list(col);width=min(48,max(14,max(len(str(c.value or '')) for c in vals[:100])+2))
                ws.column_dimensions[vals[0].column_letter].width=width
                for c in vals[1:]:
                    c.font=Font(name='Arial',size=10)
                    if isinstance(c.value,float):c.number_format='0.000000'
    os.replace(tmp,path)


def update(out, records, coverage, reference_metrics, final_tables=None):
    trials=flat_trials(records);hist=progress(records)
    csv(out/'search_history.csv',trials);csv(out/'search_progress.csv',hist)
    progress_svg(out/'search_progress.svg',records)
    sheets={'trials_validation':trials,'progress':hist,'coverage':pd.DataFrame(coverage),
            'reference_validation':pd.DataFrame(reference_metrics)}
    if records:
        best=choose_best(records)
        sheets['best_so_far']=flat_trials([best])
    if final_tables:sheets.update(final_tables)
    workbook(out/'excel/priority_weight_search.xlsx',sheets)


def prediction_frame(frame,p,threshold,audit):
    f=frame.copy().reset_index(drop=True)
    reserved={'original_row_index','csv_row_number','positive_probability','predicted_label','threshold','outcome',*ACTIVE}
    if reserved & set(f.columns):raise ValueError('Original dataset contains reserved prediction-export columns')
    f.insert(0,'original_row_index',np.arange(len(f)));f.insert(1,'csv_row_number',np.arange(len(f))+2)
    f['positive_probability']=np.asarray(p,dtype=float);f['threshold']=threshold
    f['predicted_label']=(f.positive_probability>threshold).astype(int)
    y,z=f.label.to_numpy(),f.predicted_label.to_numpy()
    f['outcome']=np.select([(y==1)&(z==1),(y==0)&(z==0),(y==0)&(z==1)],['TP','TN','FP'],default='FN')
    for c in ACTIVE:f[c]=audit[c].to_numpy(dtype=bool)
    f['canonical_smiles']=audit.canonical_smiles.to_numpy()
    f['structure_eligibility']=audit.eligibility_status.to_numpy()
    f['primary_structure_unit']=audit.primary_unit.to_numpy(dtype=bool)
    return f


def export_model(out,name,frame,p,audit):
    th={'fixed_0p5':.5,'validation_selected':choose_threshold(frame.label.to_numpy(),p)}
    summary=[]
    for policy,t in th.items():
        f=prediction_frame(frame,p,t,audit)
        folder=out/'predictions'/name/policy
        csv(folder/'val_all.csv',f)
        for outcome in ('TP','TN','FP','FN'):csv(folder/f'val_{outcome}.csv',f[f.outcome.eq(outcome)])
        summary.append(dict(experiment=name,split='val',threshold_policy=policy,**metrics(frame.label.to_numpy(),p,t)))
    return summary


def paired(frame,ref_p,best_p,audit):
    a=prediction_frame(frame,ref_p,.5,audit);b=prediction_frame(frame,best_p,.5,audit)
    f=a.drop(columns=['threshold','predicted_label','outcome','positive_probability']).copy()
    f['reference_probability']=ref_p;f['best_probability']=best_p
    f['reference_outcome']=a.outcome;f['best_outcome']=b.outcome
    f['transition']=a.outcome+'->'+b.outcome
    f['corrected_error']=a.outcome.isin(['FN','FP']) & b.outcome.isin(['TP','TN'])
    f['new_error']=a.outcome.isin(['TP','TN']) & b.outcome.isin(['FN','FP'])
    return f


def subgroup(frame,audit,p,name):
    primary=audit.primary_unit.to_numpy(dtype=bool)
    a,b=audit.V01.to_numpy(dtype=bool),audit.V03.to_numpy(dtype=bool)
    cohorts={'all_primary_723':primary,'V01':primary&a,'V03':primary&b,
             'either_rule':primary&(a|b),'both_rules':primary&a&b,'neither_rule':primary&~(a|b)}
    return [dict(experiment=name,basis='primary_isomeric_identity',cohort=c,
        **metrics(frame.label.to_numpy()[mask],p[mask],.5)) for c,mask in cohorts.items()]
