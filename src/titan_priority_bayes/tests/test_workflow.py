"""Synthetic project, actual RDKit/LightGBM, real inspect/search/resume/report I/O.

Never uses a real TITAN training cache or produces scientific performance claims.
"""
from contextlib import redirect_stdout
from pathlib import Path
import builtins
import importlib.util
import io
import json
import tempfile
import unittest
from unittest import mock

import numpy as np
import pandas as pd

from priority_bayes import cli
from priority_bayes.protocol import read_json,write_json,sha

HAS_DEPS=all(importlib.util.find_spec(n) is not None for n in ('rdkit','lightgbm','openpyxl'))


@unittest.skipUnless(HAS_DEPS,'Install actual RDKit, LightGBM and openpyxl for integration tests')
class WorkflowTests(unittest.TestCase):
    def test_real_inspect_search_resume_and_integrity_guard(self):
        from priority_bayes.training import fit_head
        from priority_bayes.report import csv
        from priority_bayes.cli import parser
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);data=root/'dataset';data.mkdir()
            src=root/'original_run';src.mkdir()
            cache=root/'original_cache';cache.mkdir()
            rules=read_json(cli.RULES_PATH)
            fv=pd.DataFrame(rules['validation_rows'],columns=['SMILES','label'])
            t=np.arange(7593)%4
            ft=pd.DataFrame({'SMILES':np.array(['CC(=O)O','CC(O)C','CCN','CCCl'])[t],
                             'label':np.array([0,1,1,0])[t]})
            frames={'train':ft,'val':fv};rng=np.random.default_rng(99)
            selected={};datasets={};X={}
            for split,f in frames.items():
                ds=data/f'{split}_classification.csv';f.to_csv(ds,index=False)
                H=np.zeros((len(f),2048),dtype=np.float32)
                H[:,:8]=rng.normal(size=(len(f),8))
                H[:,0]+=f.label.to_numpy()*1.4
                M=rng.uniform(.05,.95,size=(len(f),3)).astype(np.float32)
                hp=cache/f'{split}_H.npz';np.savez_compressed(hp,H=H)
                mp=cache/f'{split}_M.npy';np.save(mp,M)
                X[split]=np.column_stack([H,M]).astype(np.float32)
                selected[split]={'H':dict(path=str(hp),key='H',sha256=sha(hp)),
                                 'M':dict(path=str(mp),sha256=sha(mp),mechanism_order=['efflux','influx','pampa'])}
                datasets[split]=dict(path=str(ds),sha256=sha(ds))
            # Explicit forbidden paths in metadata ensure train/val scoping is
            # not merely an accident of missing held-out data.
            sentinel=data/'test_classification.csv';sentinel.write_text('DO NOT LOAD')
            datasets['test']={'path':str(sentinel),'sha256':'0'*64}
            selected['test']={'H':{'path':str(cache/'test_H.npy')},'M':{'path':str(cache/'test_M.npy')}}
            params=dict(objective='binary',metric='None',learning_rate=.08,num_leaves=7,
                max_depth=3,min_data_in_leaf=20,max_bin=31,seed=42,feature_fraction_seed=42,
                bagging_seed=42,data_random_seed=42,extra_seed=42,drop_seed=42,
                deterministic=True,force_col_wise=True,device_type='cpu',num_threads=1,
                verbosity=-1,scale_pos_weight=1.)
            model,_=fit_head(X['train'],ft.label.to_numpy(),params,30,X['val'],fv.label.to_numpy(),
                            patience=5,scale_pos_weight=1.)
            mp=cache/'original_full_fusion.txt';model.save_model(str(mp))
            p=model.predict(X['val'])
            sf=fv.copy();sf.insert(0,'original_row_index',np.arange(len(fv)))
            sf['positive_probability']=p
            csv(src/'predictions/reference_saved_full_fusion/fixed_0p5/val_all.csv',sf)
            write_json(src/'input_audit.json',dict(datasets=datasets,selected_arrays=selected,
                original_mechanism_order=['efflux','influx','pampa']))
            write_json(src/'reference_audit.json',dict(path=str(mp),sha256=sha(mp),refit_max_rounds=30,refit_patience=5))
            rules['baseline_validation']['positive_probability']=list(map(float,p))
            rules['expected_training_coverage']={'V01':dict(units=1,positive=0,negative=1,connectivity_groups=1),
                                                 'V03':dict(units=1,positive=1,negative=0,connectivity_groups=1)}
            rp=root/'synthetic_rules.json';write_json(rp,rules)
            del X,H,M,model
            output=io.StringIO()
            forbidden={'test_classification.csv','test_H.npy','test_M.npy','test_predictions.csv'}
            original_builtin,original_io=builtins.open,io.open

            def guard(original):
                def open_checked(file,*a,**kw):
                    if isinstance(file,(str,bytes,Path)) and Path(file).name in forbidden:
                        raise AssertionError(f'Forbidden test input opened: {file}')
                    return original(file,*a,**kw)
                return open_checked

            with mock.patch.object(cli,'RULES_PATH',rp),mock.patch('builtins.open',guard(original_builtin)),mock.patch('io.open',guard(original_io)),redirect_stdout(output):
                args=parser().parse_args(['--project-root',str(root),'--source-run',str(src),
                    '--startup','4','--threads','1','--budget','5','--stage','inspect'])
                out=cli.run(args)
                self.assertFalse((out/'trials').exists())
                args=parser().parse_args(['--project-root',str(root),'--resume',str(out),'--budget','5'])
                cli.run(args)
                self.assertEqual(len(list((out/'trials').glob('*/COMPLETE.json'))),5)
                self.assertTrue(read_json(out/'control_reproduction_check.json')['passed'])
                r=read_json(out/'trials/trial_0005/COMPLETE.json')
                self.assertTrue(r['proposal']['method'].startswith('GP_'))
                with mock.patch('priority_bayes.training.fit_head',side_effect=AssertionError('Completed trial unexpectedly retrained')):
                    cli.run(args)
                args=parser().parse_args(['--project-root',str(root),'--resume',str(out),'--budget','6'])
                cli.run(args)
                records=[read_json(p) for p in (out/'trials').glob('*/COMPLETE.json')]
                self.assertEqual(len({tuple(r['indices']) for r in records}),6)
                book=load_workbook(out/'excel/priority_weight_search.xlsx',read_only=True)
                self.assertIn('selected_comparison',book.sheetnames);book.close()
                p=pd.read_csv(out/'predictions/bayes_selected/fixed_0p5/val_all.csv')
                self.assertEqual(len(p),727)
                self.assertEqual(int(p.primary_structure_unit.sum()),723)
                self.assertEqual(len(pd.read_csv(out/'selected_train_weights.csv')),7593)
                best=read_json(out/'best_trial.json')
                # Cache tampering must fail before any new model fit.
                target=out/'trials'/best['trial_id']/'val_probability.npy'
                np.save(target,np.zeros(727))
                with self.assertRaisesRegex(ValueError,'artifact changed'):
                    cli.run(args)
            self.assertTrue((out/'search_progress.svg').is_file())
            self.assertEqual(parser().parse_args([]).budget,200)


if __name__=='__main__':unittest.main()
